"""
export_excel.py — экспорт данных в Excel с фильтрами и форматированием
"""
import os
from datetime import date
from openpyxl import Workbook
from openpyxl.styles import (Font, PatternFill, Alignment, Border, Side,
                              GradientFill)
from openpyxl.utils import get_column_letter
import database as db

# ── Стили ─────────────────────────────────────────────────────────────────

DARK_FILL   = PatternFill("solid", start_color="1e1e2e")
HEADER_FILL = PatternFill("solid", start_color="5b4fd4")
ODD_FILL    = PatternFill("solid", start_color="f0f0ff")
EVEN_FILL   = PatternFill("solid", start_color="ffffff")
TOTAL_FILL  = PatternFill("solid", start_color="e8e4ff")

HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
TITLE_FONT  = Font(name="Arial", bold=True, size=14, color="1e1e2e")
BODY_FONT   = Font(name="Arial", size=10)
TOTAL_FONT  = Font(name="Arial", bold=True, size=10, color="3d2db0")

BORDER_SIDE = Side(style="thin", color="cccccc")
CELL_BORDER = Border(left=BORDER_SIDE, right=BORDER_SIDE,
                     top=BORDER_SIDE, bottom=BORDER_SIDE)
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT   = Alignment(horizontal="left",   vertical="center", wrap_text=True)
RIGHT  = Alignment(horizontal="right",  vertical="center")


def _style_header(cell):
    cell.font      = HEADER_FONT
    cell.fill      = HEADER_FILL
    cell.alignment = CENTER
    cell.border    = CELL_BORDER


def _style_body(cell, even=True):
    cell.font      = BODY_FONT
    cell.fill      = EVEN_FILL if even else ODD_FILL
    cell.alignment = LEFT
    cell.border    = CELL_BORDER


def _style_total(cell):
    cell.font      = TOTAL_FONT
    cell.fill      = TOTAL_FILL
    cell.alignment = RIGHT
    cell.border    = CELL_BORDER


def _autofit(ws, min_w=10, max_w=50):
    for col in ws.columns:
        length = max(
            (len(str(cell.value)) if cell.value else 0) for cell in col
        )
        ws.column_dimensions[get_column_letter(col[0].column)].width = \
            max(min_w, min(length + 4, max_w))


def _add_title(ws, text, ncols):
    ws.merge_cells(start_row=1, start_column=1,
                   end_row=1,   end_column=ncols)
    cell = ws.cell(row=1, column=1, value=text)
    cell.font      = TITLE_FONT
    cell.alignment = CENTER
    cell.fill      = PatternFill("solid", start_color="ede9fe")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(start_row=2, start_column=1,
                   end_row=2,   end_column=ncols)
    sub = ws.cell(row=2, column=1,
                  value=f"Сформировано: {date.today().strftime('%d.%m.%Y')}")
    sub.font      = Font(name="Arial", size=9, italic=True, color="888888")
    sub.alignment = CENTER


def _write_table(ws, headers, rows, start_row=4):
    """Пишет заголовок + строки, возвращает последнюю строку."""
    for c, h in enumerate(headers, 1):
        _style_header(ws.cell(row=start_row, column=c, value=h))
    ws.row_dimensions[start_row].height = 22

    for r, row in enumerate(rows):
        even = r % 2 == 0
        for c, val in enumerate(row, 1):
            cell = ws.cell(row=start_row + 1 + r, column=c, value=val)
            _style_body(cell, even)
            if isinstance(val, (int, float)):
                cell.alignment = RIGHT
    return start_row + len(rows)


# ══════════════════════════════════════════════════════════════════
#  ЭКСПОРТ ЗАКУПОК
# ══════════════════════════════════════════════════════════════════

def export_purchases(path, date_from=None, date_to=None,
                     supplier_name=None, nomenclature_name=None,
                     warehouse_name=None):
    rows_raw = db.get_purchases()

    # Фильтрация
    rows_raw = _filter_rows(rows_raw, date_from, date_to,
                             supplier_key="supplier",
                             org_name=supplier_name,
                             nom_name=nomenclature_name,
                             wh_name=warehouse_name)

    wb = Workbook()
    ws = wb.active
    ws.title = "Закупки"
    ws.freeze_panes = "A5"

    headers = ["ID", "Дата", "Поставщик", "Номенклатура", "Ед.изм.", "Склад", "Количество"]
    _add_title(ws, "Отчёт по закупкам", len(headers))

    table_rows = [(r["id"], r["date"], r["supplier"], r["nomenclature"],
                   r["unit"], r["warehouse"], r["quantity"]) for r in rows_raw]
    last = _write_table(ws, headers, table_rows)

    # Итог
    total_row = last + 2
    ws.cell(row=total_row, column=6, value="ИТОГО:").font = TOTAL_FONT
    ws.cell(row=total_row, column=7,
            value=f"=SUM(G5:G{last})").font = TOTAL_FONT
    for c in range(1, 8):
        _style_total(ws.cell(row=total_row, column=c))

    _autofit(ws)
    wb.save(path)
    return len(table_rows)


# ══════════════════════════════════════════════════════════════════
#  ЭКСПОРТ ПРОДАЖ
# ══════════════════════════════════════════════════════════════════

def export_sales(path, date_from=None, date_to=None,
                 customer_name=None, nomenclature_name=None,
                 warehouse_name=None):
    rows_raw = db.get_sales()
    rows_raw = _filter_rows(rows_raw, date_from, date_to,
                             supplier_key="customer",
                             org_name=customer_name,
                             nom_name=nomenclature_name,
                             wh_name=warehouse_name)

    wb = Workbook()
    ws = wb.active
    ws.title = "Продажи"
    ws.freeze_panes = "A5"

    headers = ["ID", "Дата", "Покупатель", "Номенклатура", "Ед.изм.", "Склад", "Количество"]
    _add_title(ws, "Отчёт по продажам", len(headers))

    table_rows = [(r["id"], r["date"], r["customer"], r["nomenclature"],
                   r["unit"], r["warehouse"], r["quantity"]) for r in rows_raw]
    last = _write_table(ws, headers, table_rows)

    total_row = last + 2
    ws.cell(row=total_row, column=6, value="ИТОГО:").font = TOTAL_FONT
    ws.cell(row=total_row, column=7,
            value=f"=SUM(G5:G{last})").font = TOTAL_FONT
    for c in range(1, 8):
        _style_total(ws.cell(row=total_row, column=c))

    _autofit(ws)
    wb.save(path)
    return len(table_rows)


# ══════════════════════════════════════════════════════════════════
#  ЭКСПОРТ ПЕРЕМЕЩЕНИЙ
# ══════════════════════════════════════════════════════════════════

def export_movements(path, date_from=None, date_to=None,
                     nomenclature_name=None, warehouse_name=None):
    rows_raw = db.get_movements()
    rows_raw = _filter_rows(rows_raw, date_from, date_to,
                             nom_name=nomenclature_name,
                             wh_name=warehouse_name,
                             wh_keys=("from_warehouse", "to_warehouse"))

    wb = Workbook()
    ws = wb.active
    ws.title = "Перемещения"
    ws.freeze_panes = "A5"

    headers = ["ID", "Дата", "Со склада", "На склад", "Номенклатура", "Ед.изм.", "Количество"]
    _add_title(ws, "Отчёт по перемещениям", len(headers))

    table_rows = [(r["id"], r["date"], r["from_warehouse"], r["to_warehouse"],
                   r["nomenclature"], r["unit"], r["quantity"]) for r in rows_raw]
    last = _write_table(ws, headers, table_rows)

    total_row = last + 2
    ws.cell(row=total_row, column=6, value="ИТОГО:").font = TOTAL_FONT
    ws.cell(row=total_row, column=7, value=f"=SUM(G5:G{last})").font = TOTAL_FONT
    for c in range(1, 8):
        _style_total(ws.cell(row=total_row, column=c))

    _autofit(ws)
    wb.save(path)
    return len(table_rows)


# ══════════════════════════════════════════════════════════════════
#  ЭКСПОРТ ОСТАТКОВ
# ══════════════════════════════════════════════════════════════════

def export_stock(path, warehouse_name=None, nomenclature_name=None):
    whs  = db.get_all("warehouses")
    noms = db.get_all("nomenclature")
    wh_id  = next((w["id"] for w in whs  if w["name"] == warehouse_name),  None)
    nom_id = next((n["id"] for n in noms if n["name"] == nomenclature_name), None)
    rows_raw = db.get_stock_report(wh_id, nom_id)

    wb = Workbook()
    ws = wb.active
    ws.title = "Остатки"
    ws.freeze_panes = "A5"

    headers = ["Номенклатура", "Ед.изм.", "Склад", "Остаток"]
    _add_title(ws, f"Остатки на складах на {date.today().strftime('%d.%m.%Y')}", len(headers))

    table_rows = [(r["nomenclature"], r["unit"], r["warehouse"], r["balance"])
                  for r in rows_raw]
    last = _write_table(ws, headers, table_rows)

    total_row = last + 2
    ws.cell(row=total_row, column=3, value="ИТОГО ед.:").font = TOTAL_FONT
    ws.cell(row=total_row, column=4, value=f"=SUM(D5:D{last})").font = TOTAL_FONT
    for c in range(1, 5):
        _style_total(ws.cell(row=total_row, column=c))

    _autofit(ws)
    wb.save(path)
    return len(table_rows)


# ══════════════════════════════════════════════════════════════════
#  СВОДНЫЙ ОТЧЁТ (все листы)
# ══════════════════════════════════════════════════════════════════

def export_full_report(path, date_from=None, date_to=None):
    wb = Workbook()
    wb.remove(wb.active)  # удаляем дефолтный лист

    # 1. Остатки
    ws = wb.create_sheet("Остатки")
    ws.freeze_panes = "A5"
    rows = db.get_stock_report()
    headers = ["Номенклатура", "Ед.изм.", "Склад", "Остаток"]
    _add_title(ws, f"Остатки на {date.today().strftime('%d.%m.%Y')}", 4)
    table_rows = [(r["nomenclature"], r["unit"], r["warehouse"], r["balance"]) for r in rows]
    last = _write_table(ws, headers, table_rows)
    _style_total(ws.cell(row=last+2, column=3, value="ИТОГО:"))
    ws.cell(row=last+2, column=4, value=f"=SUM(D5:D{last})").font = TOTAL_FONT
    _autofit(ws)

    # 2. Закупки
    ws2 = wb.create_sheet("Закупки")
    ws2.freeze_panes = "A5"
    p_rows = _filter_rows(db.get_purchases(), date_from, date_to)
    _add_title(ws2, "Закупки", 7)
    table_rows2 = [(r["id"], r["date"], r["supplier"], r["nomenclature"],
                    r["unit"], r["warehouse"], r["quantity"]) for r in p_rows]
    last2 = _write_table(ws2, ["ID","Дата","Поставщик","Номенклатура","Ед.","Склад","Кол-во"], table_rows2)
    _style_total(ws2.cell(row=last2+2, column=6, value="ИТОГО:"))
    ws2.cell(row=last2+2, column=7, value=f"=SUM(G5:G{last2})").font = TOTAL_FONT
    _autofit(ws2)

    # 3. Продажи
    ws3 = wb.create_sheet("Продажи")
    ws3.freeze_panes = "A5"
    s_rows = _filter_rows(db.get_sales(), date_from, date_to)
    _add_title(ws3, "Продажи", 7)
    table_rows3 = [(r["id"], r["date"], r["customer"], r["nomenclature"],
                    r["unit"], r["warehouse"], r["quantity"]) for r in s_rows]
    last3 = _write_table(ws3, ["ID","Дата","Покупатель","Номенклатура","Ед.","Склад","Кол-во"], table_rows3)
    _style_total(ws3.cell(row=last3+2, column=6, value="ИТОГО:"))
    ws3.cell(row=last3+2, column=7, value=f"=SUM(G5:G{last3})").font = TOTAL_FONT
    _autofit(ws3)

    # 4. Перемещения
    ws4 = wb.create_sheet("Перемещения")
    ws4.freeze_panes = "A5"
    m_rows = _filter_rows(db.get_movements(), date_from, date_to)
    _add_title(ws4, "Перемещения", 7)
    table_rows4 = [(r["id"], r["date"], r["from_warehouse"], r["to_warehouse"],
                    r["nomenclature"], r["unit"], r["quantity"]) for r in m_rows]
    last4 = _write_table(ws4, ["ID","Дата","Со склада","На склад","Номенклатура","Ед.","Кол-во"], table_rows4)
    _style_total(ws4.cell(row=last4+2, column=6, value="ИТОГО:"))
    ws4.cell(row=last4+2, column=7, value=f"=SUM(G5:G{last4})").font = TOTAL_FONT
    _autofit(ws4)

    wb.save(path)


# ══════════════════════════════════════════════════════════════════
#  ВСПОМОГАТЕЛЬНЫЕ ФИЛЬТРЫ
# ══════════════════════════════════════════════════════════════════

def _filter_rows(rows, date_from=None, date_to=None,
                 supplier_key=None, org_name=None,
                 nom_name=None, wh_name=None,
                 wh_keys=("warehouse",)):
    result = []
    for r in rows:
        if date_from and str(r.get("date","")) < date_from:
            continue
        if date_to   and str(r.get("date","")) > date_to:
            continue
        if org_name and supplier_key:
            if r.get(supplier_key,"") != org_name:
                continue
        if nom_name and r.get("nomenclature","") != nom_name:
            continue
        if wh_name:
            match = any(r.get(k,"") == wh_name for k in wh_keys)
            if not match:
                continue
        result.append(r)
    return result
